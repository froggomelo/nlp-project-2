"""Blind check of the kept test pairs that the floor is computed on (no retraining, CPU, seconds).

A random sample of kept test pairs, equal per similarity quartile, is split between three labellers;
30 rows are labelled twice so agreement can be reported. Each labeller gets their own file with the
similarity hidden, so neither the score nor anyone else's label can sway them.

Usage:
    python src/verify_pairs.py sheet      # -> labels/verify/pairs_verify_<name>.csv + pairs_verify_key.csv
    python src/verify_pairs.py xlsx       # same sheets as workbooks: Labels, isiXhosa, English, Side by side tabs
    python src/verify_pairs.py part2 --name Gopolang   # the pairs not in that person's sheet, for one labeller
    python src/verify_pairs.py score      # precision, Cohen's kappa, floor and gap on verified pairs
                                          # -> results/raw/pair_verification.json, and every label used
                                          #    -> labels/verify/pairs_verified.csv
    python src/verify_pairs.py score --frozen   # rescore from pairs_verified.csv (no workbooks needed)

Labels (labels/LABELLING_GUIDE.md, Job 1b): 2 = translation, 1 = partial, 0 = different, blank = can't tell.
"""
import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
from sklearn.metrics import cohen_kappa_score

LABELLERS = ["Gopolang", "Tawananyasha", "Lily"]
OUT = Path("labels/verify")
KEY = OUT / "pairs_verify_key.csv"
FROZEN = OUT / "pairs_verified.csv"  # every label the last score used, one row per sampled pair
B = 1000


def person_file(name):
    return OUT / f"pairs_verify_{name}.csv"


def person_xlsx(name):
    return OUT / f"pairs_verify_{name}.xlsx"


def part2_file(name):
    return OUT / f"pairs_verify_{name}_part2.csv"


def part2_xlsx(name):
    return OUT / f"pairs_verify_{name}_part2.xlsx"


def read_one(xlsx, csv):
    if xlsx.exists():
        d, src = pd.read_excel(xlsx, sheet_name="Labels", dtype={"label": "string"}), xlsx
        if d["label"].isna().all():  # some labellers typed labels next to the texts on the Side by side tab
            side = pd.read_excel(xlsx, sheet_name="Side by side", dtype=str)
            extra = [c for c in side.columns if c not in ("row_id", "isiXhosa", "English")]
            if extra and side[extra[0]].notna().any():
                d = side[["row_id", extra[0]]].rename(columns={extra[0]: "label"})
                src = f"{xlsx} (Side by side tab, column '{extra[0]}')"
    else:
        d, src = pd.read_csv(csv, dtype={"label": "string"}), csv
    print(f"labels from {src}")
    return d.set_index("row_id")["label"]


def read_labels(name):
    """Labels from the workbook's Labels tab if the workbook exists, else from the CSV; plus any part 2."""
    lab = read_one(person_xlsx(name), person_file(name))
    if part2_xlsx(name).exists() or part2_file(name).exists():
        lab = pd.concat([lab, read_one(part2_xlsx(name), part2_file(name))])
    return lab


def write_xlsx(d, out):
    """Labels tab to fill in, the texts on their own tabs, and both side by side."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    tabs = {"Labels": d[["row_id", "shared", "label", "note"]],
            "isiXhosa": d[["row_id", "xh_text"]].rename(columns={"xh_text": "isiXhosa"}),
            "English": d[["row_id", "en_text"]].rename(columns={"en_text": "English"}),
            "Side by side": d[["row_id", "xh_text", "en_text"]].rename(
                columns={"xh_text": "isiXhosa", "en_text": "English"})}
    widths = {"row_id": 8, "shared": 40, "label": 8, "note": 50, "isiXhosa": 90, "English": 90}
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for tab, t in tabs.items():
            t.to_excel(xw, sheet_name=tab, index=False)
            ws = xw.sheets[tab]
            ws.freeze_panes = "B2"
            for j, col in enumerate(t.columns):
                letter = ws.cell(row=1, column=j + 1).column_letter
                ws.column_dimensions[letter].width = widths[col]
                ws.cell(row=1, column=j + 1).font = Font(bold=True)
                for r in range(2, len(t) + 2):
                    ws.cell(row=r, column=j + 1).alignment = Alignment(wrap_text=True, vertical="top")
            if tab == "Labels":
                dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"C2:C{len(t) + 1}")


def make_xlsx(args):
    """One workbook per labeller, built from the CSVs, so any labels already typed there are carried over."""
    for name in LABELLERS:
        out = person_xlsx(name)
        if out.exists() and not args.force:
            print(f"{name}: {out} exists, skipped (use --force to rebuild from the CSV)")
            continue
        d = pd.read_csv(person_file(name), dtype=str).fillna("")
        write_xlsx(d, out)
        print(f"{name}: {len(d)} rows -> {out}")


def make_part2(args):
    """For one labeller covering everyone: the sampled pairs not already in their sheet, as one workbook."""
    if part2_xlsx(args.name).exists() and not args.force:
        raise SystemExit(f"{part2_xlsx(args.name)} exists; use --force to overwrite")
    mine = set(pd.read_csv(person_file(args.name), dtype=str)["row_id"])
    rest = pd.concat([pd.read_csv(person_file(n), dtype=str) for n in LABELLERS if n != args.name])
    rest = rest[~rest.row_id.isin(mine)].drop_duplicates("row_id")
    rest = rest.sample(frac=1, random_state=args.seed + 7).fillna("").assign(label="", note="")
    rest.to_csv(part2_file(args.name), index=False, encoding="utf-8-sig")
    write_xlsx(rest, part2_xlsx(args.name))
    print(f"{args.name}: {len(rest)} more rows -> {part2_xlsx(args.name)}")


def make_sheet(args):
    started = [p for p in map(person_file, LABELLERS)
               if p.exists() and pd.read_csv(p, dtype=str)["label"].notna().any()]
    if started and not args.force:
        raise SystemExit(f"labels already entered in {[str(p) for p in started]}; use --force to overwrite")
    cand = pd.read_parquet("results/raw/pair_candidates.parquet")
    kept = cand[(cand.split == "test") & (cand.sim >= args.threshold)].copy()
    kept["quartile"] = pd.qcut(kept.sim, 4, labels=[1, 2, 3, 4]).astype(int)
    per_q = args.n // 4
    sample = kept.groupby("quartile").sample(n=per_q, random_state=args.seed)
    sample = sample.sample(frac=1, random_state=args.seed + 1).reset_index(drop=True)  # ids hide the quartile
    sample["row_id"] = [f"P{i + 1:03d}" for i in range(len(sample))]

    # three blocks of main labels; 10 rows of each block go to the next person as second labels
    blocks = np.array_split(np.arange(len(sample)), len(LABELLERS))
    sample["main"], sample["second"] = "", ""
    for j, idx in enumerate(blocks):
        sample.loc[idx, "main"] = LABELLERS[j]
        sample.loc[idx[: args.overlap // len(LABELLERS)], "second"] = LABELLERS[(j + 1) % len(LABELLERS)]

    OUT.mkdir(parents=True, exist_ok=True)
    sample[["row_id", "issue", "year", "xh_id", "en_id", "sim", "quartile", "main", "second"]].to_csv(
        KEY, index=False, encoding="utf-8-sig")
    for j, name in enumerate(LABELLERS):
        mine = sample[(sample["main"] == name) | (sample["second"] == name)]
        mine = mine.sample(frac=1, random_state=args.seed + 2 + j)  # second-label rows not identifiable
        sheet = mine[["row_id", "issue", "xh_text", "en_text", "shared"]].assign(label="", note="")
        sheet.to_csv(person_file(name), index=False, encoding="utf-8-sig")
        print(f"{name}: {len(sheet)} rows -> {person_file(name)}")
    print(f"{len(sample)} pairs from {len(kept)} kept test pairs, {per_q} per similarity quartile; key -> {KEY}")


def ba(df):
    return 0.5 * ((df.loc[df.y == 1, "pred"] == 1).mean() + (df.loc[df.y == 0, "pred"] == 0).mean())


def boot(pred, fn, rng):
    issues = pred["issue"].unique()
    groups = {i: g for i, g in pred.groupby("issue")}
    vals = []
    for _ in range(B):
        v = fn(pd.concat([groups[i] for i in rng.choice(issues, len(issues))]))
        if np.isfinite(v):
            vals.append(v)
    return [float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5))]


def score(args):
    key = pd.read_csv(KEY)
    if args.frozen:  # rescore from the saved label table instead of the workbooks
        frozen = pd.read_csv(FROZEN, dtype={"label": "float", "label_2": "float", "by": "string"})
        key = key.merge(frozen[["row_id", "label", "by", "label_2"]], on="row_id", how="left")
        key["by"] = key["by"].fillna("")
        print(f"labels from {FROZEN}")
        return report(key, args)
    labels = {name: read_labels(name) for name in LABELLERS}

    def lab(row_id, who):
        v = labels[who].get(row_id)
        v = None if v is None or pd.isna(v) else str(v).strip().removesuffix(".0")
        if v not in (None, "", "0", "1", "2"):
            raise SystemExit(f"{row_id} ({who}): label must be 0, 1, 2 or blank, got {v!r}")
        return float(v) if v else np.nan

    # a row's label comes from its main labeller if they gave one, else from anyone who did;
    # label_2 is a second person's label for the same row, used only for agreement
    def pick(row):
        order = [row["main"]] + [n for n in [row["second"]] + LABELLERS if isinstance(n, str) and n != row["main"]]
        got = [(n, lab(row["row_id"], n)) for n in dict.fromkeys(order)]
        got = [(n, v) for n, v in got if not np.isnan(v)]
        return pd.Series({"label": got[0][1] if got else np.nan, "by": got[0][0] if got else "",
                          "label_2": got[1][1] if len(got) > 1 else np.nan})

    key = key.join(key.apply(pick, axis=1))
    # one plain table of every label used, so the scores can be rerun and checked without the workbooks
    key[["row_id", "issue", "xh_id", "en_id", "sim", "quartile", "label", "by", "label_2"]].to_csv(
        FROZEN, index=False, encoding="utf-8")
    print(f"label table -> {FROZEN}")
    return report(key, args)


def report(key, args):
    done = key[key.label.notna()]
    print(f"labelled {len(done)} of {len(key)} pairs ({key.label.isna().sum()} blank); by {done.by.value_counts().to_dict()}")

    def shares(d):
        return {"n": int(len(d)), "translation": round(float((d.label == 2).mean()), 3),
                "partial": round(float((d.label == 1).mean()), 3), "different": round(float((d.label == 0).mean()), 3)}

    out = {"overall": shares(done), "by_quartile": {int(q): shares(g) for q, g in done.groupby("quartile")}}
    print("overall", out["overall"])
    for q, s in out["by_quartile"].items():
        print(f"  similarity quartile {q}: {s}")

    both = key.dropna(subset=["label", "label_2"])
    if len(both) >= 2:
        out["agreement"] = {"n": int(len(both)), "raw": round(float((both.label == both.label_2).mean()), 3),
                            "cohen_kappa": round(float(cohen_kappa_score(both.label, both.label_2)), 3)}
        print("agreement", out["agreement"])

    # AI-proposed labels (labels/verify/pairs_verify_AI_*.csv) are never used as labels above; they are
    # only compared with the human labels on the same rows, so the report can say how often they agree
    ai_files = sorted(OUT.glob("pairs_verify_AI_*.csv"))
    if ai_files:
        ai = pd.concat([pd.read_csv(f, dtype={"label": "string"}) for f in ai_files]).set_index("row_id")["label"]
        ai = pd.to_numeric(ai, errors="coerce").dropna()
        cmp = done.set_index("row_id").label.reindex(ai.index).dropna()
        out["ai_vs_human"] = {"ai_rows": int(len(ai)), "with_human_label": int(len(cmp))}
        if len(cmp) >= 2:
            a = ai.loc[cmp.index]
            out["ai_vs_human"].update(raw=round(float((a == cmp).mean()), 3),
                                      cohen_kappa=round(float(cohen_kappa_score(cmp, a)), 3))
        print("AI labels vs human (not used as labels):", out["ai_vs_human"])

    # floor and gap on verified pairs, from the saved test predictions (no retraining)
    pred = pd.read_parquet("results/raw/stage1_test_predictions.parquet")
    rng = np.random.default_rng(13)
    for name, keep in [("verified (label 2)", done.label == 2), ("label 1 or 2", done.label >= 1),
                       ("all sampled", done.label.notna())]:
        ids = set(done.loc[keep, "xh_id"]) | set(done.loc[keep, "en_id"])
        def floor(d):
            return ba(d[d.chunk_id.isin(ids)])
        def gap(d):
            return ba(d) - ba(d[d.chunk_id.isin(ids)])
        out[name] = {"pairs": int(keep.sum()), "floor": round(float(floor(pred)), 3),
                     "floor_ci": [round(v, 3) for v in boot(pred, floor, rng)],
                     "gap": round(float(gap(pred)), 3), "gap_ci": [round(v, 3) for v in boot(pred, gap, rng)]}
        print(f"{name}: {out[name]}")

    path = Path("results/raw/pair_verification.json")
    path.write_text(json.dumps(out, indent=1))
    print(f"-> {path}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mode", choices=["sheet", "xlsx", "part2", "score"])
    ap.add_argument("--name", default="Gopolang", help="part2: the labeller who will cover the other sheets' pairs")
    ap.add_argument("--n", type=int, default=100)
    ap.add_argument("--overlap", type=int, default=30)
    ap.add_argument("--threshold", type=float, default=0.55)
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--frozen", action="store_true",
                    help="score: read labels from labels/verify/pairs_verified.csv instead of the workbooks")
    args = ap.parse_args()
    {"sheet": make_sheet, "xlsx": make_xlsx, "part2": make_part2, "score": score}[args.mode](args)


if __name__ == "__main__":
    main()
